from io import BytesIO
from uuid import uuid4
from unittest.mock import patch

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import Client, User
from apps.core.models import SectionNote, SectionNoteEvent
from apps.core.release_scope import release_scope_context
from apps.core.section_notes import scope


class ReleaseScopeNotesTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.planner = User.objects.create_user(username="notes-planner", role=User.Role.PLANNING)
        cls.viewer = User.objects.create_user(username="notes-viewer", role=User.Role.VIEWER)
        cls.records = {}
        for label, context in (
            ("legacy", {}), ("official", {"release_scope": "official"}),
            ("preparing", {"release_scope": "preparing"}), ("both", {"release_scope": "all"}),
        ):
            cls.records[label] = SectionNote.objects.create(
                author=cls.planner, author_name="Notes planner", section="s05", panel="model-review",
                note_date=timezone.localdate(), body=f"{label} private note", information_only=True, context=context,
            )

    def setUp(self):
        self.client.force_login(self.planner)
        self.url = reverse("core:section_notes")
        network = patch("psycopg2.connect", side_effect=AssertionError("No operational sources for notes tests"))
        network.start()
        self.addCleanup(network.stop)

    def data(self, **changes):
        return {
            "section": "s05", "panel": "model-review", "date": timezone.localdate().isoformat(),
            "body": "A new note", "information_only": True, "request_id": str(uuid4()), **changes,
        }

    def create(self, selected, data):
        return self.client.post(self.url + f"?release_scope={selected}", data, content_type="application/json")

    def change(self, selected, note, **changes):
        return self.client.post(
            reverse("core:section_note_status", args=[note.pk]) + f"?release_scope={selected}",
            {"status": "resolved", "version": note.version, **changes}, content_type="application/json",
        )

    def test_list_and_summary_keep_legacy_official_preparing_and_both_separate(self):
        cases = (
            (self.viewer, "preparing", {"legacy", "official"}),
            (self.planner, "official", {"legacy", "official"}),
            (self.planner, "preparing", {"preparing"}),
            (self.planner, "all", {"legacy", "official", "preparing", "both"}),
        )
        for user, selected, expected in cases:
            with self.subTest(role=user.role, selected=selected):
                self.client.force_login(user)
                response = self.client.get(self.url, {"panel": "model-review", "release_scope": selected})
                self.assertEqual(response.status_code, 200)
                notes = response.json()["notes"]
                self.assertEqual({note["id"] for note in notes}, {self.records[label].pk for label in expected})
                summary = self.client.get(self.url, {"release_scope": selected}).json()["sections"]["model-review"]
                self.assertEqual(summary["total"], len(expected))
                if "legacy" in expected:
                    legacy = next(note for note in notes if note["id"] == self.records["legacy"].pk)
                    self.assertEqual(legacy["context"]["release_scope"], "official")

    def test_server_stamps_selected_scope_and_ignores_forged_context(self):
        cases = ((self.viewer, "preparing", "official"), (self.planner, "preparing", "preparing"), (self.planner, "all", "all"))
        for user, selected, expected in cases:
            with self.subTest(role=user.role, selected=selected):
                self.client.force_login(user)
                response = self.create(selected, self.data(context={"release_scope": "forged", "discipline": "piping", "secret": "discard"}))
                self.assertEqual(response.status_code, 201)
                returned = response.json()["note"]
                self.assertEqual(returned["context"], {"discipline": "piping", "release_scope": expected})
                self.assertEqual(SectionNote.objects.get(pk=returned["id"]).context, returned["context"])

    def test_status_changes_require_visibility_and_keep_audit_for_authorized_both(self):
        preparing = self.records["preparing"]
        self.client.force_login(self.viewer)
        self.assertEqual(self.change("preparing", preparing).status_code, 404)
        self.assertEqual(self.change("all", preparing).status_code, 404)
        self.assertFalse(SectionNoteEvent.objects.filter(note=preparing).exists())
        self.client.force_login(self.planner)
        self.assertEqual(self.change("official", preparing).status_code, 404)
        response = self.change("preparing", preparing)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["note"]["events"][-1]["to"], "resolved")
        official = self.records["official"]
        self.assertEqual(self.change("all", official).status_code, 200)
        self.assertEqual(self.change("preparing", self.records["both"]).status_code, 404)
        self.assertEqual(self.change("all", self.records["both"]).status_code, 200)

    def test_minutes_export_contains_only_visible_notes_and_identifies_selected_scope(self):
        cases = (
            (self.viewer, "all", {"legacy", "official"}, "Official"),
            (self.planner, "preparing", {"preparing"}, "Preparing"),
            (self.planner, "all", {"legacy", "official", "preparing", "both"}, "Both"),
        )
        for user, selected, expected, label in cases:
            with self.subTest(role=user.role, selected=selected):
                self.client.force_login(user)
                response = self.client.get(reverse("core:section_minutes"), {"release_scope": selected})
                self.assertEqual(response.status_code, 200)
                book = load_workbook(BytesIO(response.content), read_only=True)
                try:
                    rows = list(book.active.iter_rows(min_row=6, values_only=True))
                    self.assertEqual({row[5] for row in rows}, {self.records[key].body for key in expected})
                    self.assertIn(f"Drawings: {label}", book.active["A2"].value)
                    if "legacy" in expected:
                        legacy_row = next(row for row in rows if row[0] == self.records["legacy"].pk)
                        self.assertIn("release_scope: official", legacy_row[13])
                finally:
                    book.close()

    def test_uuid_retries_from_another_scope_never_return_or_move_private_note(self):
        data = self.data()
        first = self.create("preparing", data)
        self.assertEqual(first.status_code, 201)
        private_id = first.json()["note"]["id"]
        self.assertEqual(self.create("preparing", data).status_code, 200)
        for selected in ("official", "all"):
            with self.subTest(selected=selected):
                response = self.create(selected, data)
                self.assertEqual(response.status_code, 409)
                self.assertNotIn("note", response.json())
        note = SectionNote.objects.get(pk=private_id)
        self.assertEqual(note.context["release_scope"], "preparing")
        self.assertEqual(note.events.count(), 1)
        User.objects.filter(pk=self.planner.pk).update(role=User.Role.VIEWER)
        self.assertEqual(self.create("preparing", data).status_code, 409)

    def test_legacy_uuid_retry_remains_official_without_rewriting_history(self):
        legacy = self.records["legacy"]
        data = self.data(request_id=str(legacy.request_id), body=legacy.body)
        response = self.create("official", data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["note"]["context"]["release_scope"], "official")
        legacy.refresh_from_db()
        self.assertEqual(legacy.context, {})
        self.assertEqual(self.create("preparing", data).status_code, 409)

    def test_explicit_tab_scope_survives_another_request_changing_the_session(self):
        self.client.get(self.url, {"release_scope": "preparing"})
        self.client.get(self.url, {"release_scope": "official"})
        original_tab = self.client.get(self.url, {"release_scope": "preparing", "panel": "model-review"})
        self.assertEqual([note["id"] for note in original_tab.json()["notes"]], [self.records["preparing"].pk])
        self.assertEqual(self.change("preparing", self.records["preparing"]).status_code, 200)

    def test_scope_helper_clamps_viewer_even_when_called_in_privileged_context(self):
        with release_scope_context("all"):
            self.assertEqual(set(scope(self.viewer).values_list("pk", flat=True)), {
                self.records["legacy"].pk, self.records["official"].pk,
            })

    def test_tenant_filter_still_applies_in_both_scope(self):
        tenant = Client.objects.create(name="Isolated notes client", slug="isolated-notes")
        user = User.objects.create_user(username="tenant-supply-notes", role=User.Role.SUPPLY, client=tenant)
        self.client.force_login(user)
        response = self.client.get(self.url, {"release_scope": "all", "panel": "model-review"})
        self.assertEqual(response.json()["total"], 0)
        self.assertEqual(self.change("all", self.records["preparing"]).status_code, 404)

    def test_modal_identifies_the_scope_bound_to_the_current_page(self):
        response = self.client.get(reverse("core:home"), {"release_scope": "preparing"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-release-scope="preparing"')
        self.assertContains(response, "PANEL FOLLOW-UP · Preparing")
