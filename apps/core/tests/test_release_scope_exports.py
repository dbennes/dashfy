"""DATAFY exports are scoped; independent engineering registers stay whole."""
from io import BytesIO
import json
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.core.datafy_supply import supply_filters_hash
from apps.core.models import (
    DatafySupplySnapshot, EngineeringDocumentStatus,
    EngineeringMonitorImport, EngineeringStatusImport,
)
from apps.core.release_scope import current_release_scope


class ReleaseScopeExportTests(TestCase):
    OFFICIAL = "EXPORT-OFFICIAL-001"
    PREPARING = "EXPORT-PREPARING-001"

    @classmethod
    def setUpTestData(cls):
        cls.viewer = User.objects.create_user(username="export-viewer", role=User.Role.VIEWER)
        cls.planner = User.objects.create_user(username="export-planner", role=User.Role.PLANNING)
        cls.admin = User.objects.create_user(username="export-admin", role=User.Role.ADMIN)
        cls.monitor = EngineeringMonitorImport.objects.create(
            original_filename="scoped-monitor.xlsx", file_hash="scope-monitor",
            document_count=2, monitored_document_count=1, excluded_count=1,
            discipline_count=2, metadata={"import_mode": "mdr_engineering"},
            payload={"documents": [
                {
                    "document_number": cls.OFFICIAL, "title": "Official drawing",
                    "row_number": 2, "discipline": "PIPING", "source_discipline": "PIPING",
                    "revision": "C01", "is_monitored": True, "is_countable": True,
                    "status_bucket": "AFC 1",
                },
                {
                    "document_number": cls.PREPARING, "title": "Preparing drawing",
                    "row_number": 3, "discipline": "STRUCTURAL", "source_discipline": "STRUCTURAL",
                    "revision": "C02", "is_monitored": True, "is_countable": False,
                    "status_bucket": "IFR", "excluded_reason": "Excluded title",
                },
            ]},
        )
        cls.status = EngineeringStatusImport.objects.create(
            original_filename="scoped-DED.xlsx", file_hash="scope-ded", document_count=2, discipline_count=2,
        )
        EngineeringDocumentStatus.objects.create(
            import_batch=cls.status, row_number=2, document_number=cls.OFFICIAL,
            title="Official DED drawing", discipline="PIPING", revision="C01",
            doc_status_group="AFC", afc_code="1",
        )
        EngineeringDocumentStatus.objects.create(
            import_batch=cls.status, row_number=3, document_number=cls.PREPARING,
            title="Preparing DED drawing", discipline="STRUCTURAL", revision="C02",
            doc_status_group="IFR",
        )

    def setUp(self):
        network = patch("psycopg2.connect", side_effect=AssertionError("Operational database access is forbidden"))
        network.start()
        self.addCleanup(network.stop)

    def workbook(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        self.addCleanup(workbook.close)
        return workbook

    def cases(self):
        return (
            (self.viewer, "official", {self.OFFICIAL, self.PREPARING}),
            (self.viewer, "preparing", {self.OFFICIAL, self.PREPARING}),
            (self.planner, "official", {self.OFFICIAL, self.PREPARING}),
            (self.planner, "preparing", {self.OFFICIAL, self.PREPARING}),
            (self.planner, "all", {self.OFFICIAL, self.PREPARING}),
        )

    def test_monitor_export_and_summary_are_independent_of_datafy_scope(self):
        for user, requested, expected in self.cases():
            with self.subTest(role=user.role, requested=requested):
                self.client.force_login(user)
                response = self.client.get(reverse("core:export_engineering_monitor"), {"release_scope": requested})
                workbook = self.workbook(response)
                rows = list(workbook["Base Engenharia"].iter_rows(min_row=2, values_only=True))
                self.assertEqual({row[0] for row in rows}, expected)
                summary = {
                    row[0]: row[1] for row in workbook["Resumo"].iter_rows(values_only=True)
                    if row[0] in {"Linhas brutas", "Contabilizadas", "Desconsideradas"}
                }
                expected_countable = int(self.OFFICIAL in expected)
                self.assertEqual(summary, {
                    "Linhas brutas": len(expected),
                    "Contabilizadas": expected_countable,
                    "Desconsideradas": len(expected) - expected_countable,
                })
                self.assertEqual(sum(row[5] == "SIM" for row in rows), expected_countable)

    def test_ded_export_and_totals_are_independent_of_datafy_scope(self):
        for user, requested, expected in self.cases():
            with self.subTest(role=user.role, requested=requested):
                self.client.force_login(user)
                workbook = self.workbook(self.client.get(
                    reverse("core:export_engineering_status"), {"release_scope": requested},
                ))
                rows = list(workbook["Drawings"].iter_rows(min_row=2, values_only=True))
                self.assertEqual({row[0] for row in rows}, expected)
                statuses = {row[0]: row[1] for row in workbook["Status"].iter_rows(min_row=2, values_only=True)}
                disciplines = {row[0]: row[1] for row in workbook["Disciplinas"].iter_rows(min_row=2, values_only=True)}
                expected_statuses = {}
                expected_disciplines = {}
                if self.OFFICIAL in expected:
                    expected_statuses["AFC 1"] = 1
                    expected_disciplines["PIPING"] = 1
                if self.PREPARING in expected:
                    expected_statuses["IFR"] = 1
                    expected_disciplines["STRUCTURAL"] = 1
                self.assertEqual(statuses, expected_statuses)
                self.assertEqual(disciplines, expected_disciplines)
                self.assertEqual(sum(statuses.values()), len(rows))
                self.assertEqual(sum(disciplines.values()), len(rows))

    def test_engineering_exports_work_while_datafy_is_unavailable(self):
        self.client.force_login(self.viewer)
        for endpoint in ("core:export_engineering_monitor", "core:export_engineering_status"):
            with self.subTest(endpoint=endpoint), patch(
                "apps.core.real_sources._datafy_conn", side_effect=AssertionError("DATAFY unavailable"),
            ) as datafy:
                response = self.client.get(reverse(endpoint), {"release_scope": "all"})
                self.assertEqual(response.status_code, 200)
                self.assertIn("Content-Disposition", response)
                datafy.assert_not_called()

    def test_viewer_supply_export_cannot_forge_preparing_filters_or_metadata(self):
        self.client.force_login(self.viewer)
        captured = []

        def source(filters):
            captured.append(dict(filters))
            scope = current_release_scope()
            return {
                "source_mode": "postgres_live", "source_database": "fixture-datafy",
                "drawing_release_scope": f"{scope}-v2", "material_flow": {"total": 1},
                "supply_campaign_views": [{"totals": {"drawings": 1}}],
                "material_rows": [{"drawing_number": self.OFFICIAL, "material_item_id": 123}],
            }

        with patch("apps.core.views.real_sources._construction_datafy_snapshot", side_effect=source):
            response = self.client.get(reverse("core:export_datafy_supply"), {"release_scope": "preparing"})
        self.assertEqual(response.status_code, 200)
        snapshot = response.json()["snapshot"]
        self.assertEqual(captured[0]["release_scope"], "official")
        self.assertEqual(snapshot["filters"]["release_scope"], "official")
        self.assertEqual(snapshot["filters"]["release_scope_version"], 1)
        self.assertEqual(snapshot["payload"]["drawing_release_scope"], "official-v2")
        self.assertEqual(snapshot["filters_hash"], supply_filters_hash(snapshot["filters"]))
        self.assertEqual(snapshot["payload"]["material_rows"][0]["drawing_number"], self.OFFICIAL)
        self.assertEqual(response["X-Drawing-Release-Scope"], "official")

    def test_admin_import_with_mismatched_scope_or_token_preserves_existing_snapshots(self):
        self.client.force_login(self.admin)
        for scope in ("official", "preparing"):
            filters = {"release_scope": scope, "release_scope_version": 1}
            DatafySupplySnapshot.objects.create(
                filters=filters, filters_hash=supply_filters_hash(filters),
                payload={"drawing_release_scope": f"{scope}-v2", "original": True},
            )
        before = list(DatafySupplySnapshot.objects.order_by("pk").values())
        cases = (
            ("official", {"release_scope": "preparing", "release_scope_version": 1}, "preparing-v2"),
            ("preparing", {"release_scope": "preparing", "release_scope_version": 1}, "official-v2"),
            ("official", {"release_scope": "official", "release_scope_version": 1}, None),
            ("official", {"release_scope": "official"}, "official-v2"),
        )
        for selected, filters, token in cases:
            with self.subTest(selected=selected, filters=filters, token=token):
                payload = {"material_flow": {"total": 99}, "material_rows": [{"private": "do not store"}]}
                if token is not None:
                    payload["drawing_release_scope"] = token
                upload = SimpleUploadedFile(
                    "mismatched-supply.json", json.dumps({"snapshot": {"filters": filters, "payload": payload}}).encode(),
                    content_type="application/json",
                )
                response = self.client.post(
                    reverse("core:import_datafy_supply") + f"?release_scope={selected}", {"datafy_supply_file": upload},
                )
                self.assertEqual(response.status_code, 302)
                self.assertEqual(list(DatafySupplySnapshot.objects.order_by("pk").values()), before)
                self.assertEqual(DatafySupplySnapshot.objects.filter(is_active=True).count(), 2)
