from copy import deepcopy
from io import BytesIO
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.core import rundown_workbook as wb
from apps.core.models import RundownImport
from .test_rundown_modes_ui import rundown_modes_fixture


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class RundownWorkbookTests(TestCase):
    def setUp(self):
        self.modes = rundown_modes_fixture()
        self.enterContext(patch.object(wb, "source_modes", side_effect=lambda: deepcopy(self.modes)))
        self.admin = User.objects.create_user(username="rundown-admin", role="admin")
        self.client.force_login(self.admin)
        self.current = wb.current_state()
        self.content = wb.export_workbook(self.current)

    def edit(self, callback):
        book = load_workbook(BytesIO(self.content))
        callback(book)
        stream = BytesIO()
        book.save(stream)
        book.close()
        return stream.getvalue()

    def test_round_trip_and_empty_sample_templates(self):
        parsed = wb.parse_workbook(self.content, self.current)
        self.assertEqual(parsed["updates"], {})
        book = load_workbook(BytesIO(self.content), data_only=True)
        self.assertEqual(len(book.sheetnames), 8)
        for title in wb.SHEETS.values():
            self.assertIn(title, book.sheetnames)
        sample = book["Piping - Installation"]
        self.assertEqual(sample["B5"].value, "NO")
        self.assertIsNone(sample["A10"].value)
        self.assertEqual(book["Piping - Fabrication"]["E10"].value, 10)
        book.close()

    def test_progress_import_is_persistent_and_does_not_change_other_views(self):
        content = self.edit(lambda book: setattr(book["Piping - Fabrication"]["D10"], "value", 3))
        parsed = wb.parse_workbook(content, self.current)
        batch = wb.apply_import(parsed, "updated.xlsx", "abc", self.admin)
        self.assertEqual(set(batch.payload), {"fabrication:piping"})
        result = wb.overlay_modes(self.modes)
        piping = result["fabrication"]["disciplines"]["piping"]
        self.assertEqual(piping["kpis"]["actual_completed"], 3)
        self.assertEqual(piping["kpis"]["actual_progress_pct"], 30)
        self.assertEqual(piping["charts"]["actual_rundown"][0], 7)
        self.assertIsNone(piping["charts"]["actual_rundown"][1])
        self.assertEqual(result["installation"], self.modes["installation"])
        current = wb.current_state()
        self.assertEqual(wb.parse_workbook(wb.export_workbook(current), current)["updates"], {})
        with self.assertRaisesRegex(ValueError, "changed"):
            wb.apply_import(parsed, "stale.xlsx", "abc", self.admin)
        self.assertEqual(RundownImport.objects.count(), 1)

    def test_real_installation_replaces_only_its_sample(self):
        def change(book):
            sheet = book["Electrical - Installation"]
            for cell, value in {"B2": 20, "B3": "2026-09-15", "B5": "YES", "A10": "2026-09-14", "B10": 20, "D10": 5}.items():
                sheet[cell] = value
        parsed = wb.parse_workbook(self.edit(change), self.current)
        wb.apply_import(parsed, "real.xlsx", "abc", self.admin)
        result = wb.overlay_modes(self.modes)
        electrical = result["installation"]["disciplines"]["electrical"]
        self.assertFalse(electrical["source"]["is_sample"])
        self.assertEqual(electrical["kpis"]["actual_progress_pct"], 25)
        self.assertTrue(result["installation"]["disciplines"]["piping"]["source"]["is_sample"])

    def test_rejects_invalid_quantities_dates_formulas_duplicates_and_units(self):
        for cell, value, message in [("D10", 11, "exceeds"), ("D10", -1, "non-negative"),
                                     ("D10", "=1+1", "whole number"), ("D10", 0.5, "whole number"),
                                     ("A10", "invalid", "valid date"), ("A11", "2026-09-01", "duplicate"),
                                     ("B4", "tonnes", "unit")]:
            with self.subTest(cell=cell, value=value):
                content = self.edit(lambda book: setattr(book["Piping - Fabrication"][cell], "value", value))
                with self.assertRaisesRegex(ValueError, message):
                    wb.parse_workbook(content, self.current)
        def future(book):
            sheet = book["Piping - Fabrication"]
            sheet["A11"] = "2026-09-16"
            sheet["D11"] = 1
        with self.assertRaisesRegex(ValueError, "after Data date"):
            wb.parse_workbook(self.edit(future), self.current)
        self.assertFalse(RundownImport.objects.exists())

    def test_missing_sheet_and_stale_source_are_rejected(self):
        content = self.edit(lambda book: book.remove(book["Electrical - Fabrication"]))
        with self.assertRaisesRegex(ValueError, "Missing sheet"):
            wb.parse_workbook(content, self.current)
        self.modes["fabrication"]["disciplines"]["piping"]["kpis"]["scope_total"] += 1
        with self.assertRaisesRegex(ValueError, "changed since export"):
            wb.parse_workbook(self.content, wb.current_state())

    def test_preview_apply_permissions_and_user_bound_token(self):
        url = reverse("core:import_rundown")
        exported = self.client.get(reverse("core:export_rundown"))
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(exported["Cache-Control"], "no-store")
        content = self.edit(lambda book: setattr(book["Piping - Fabrication"]["D10"], "value", 2))
        response = self.client.post(url, {"rundown_file": SimpleUploadedFile("updated.xlsx", content)})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(RundownImport.objects.exists())
        token = response.context["preview_token"]
        other = User.objects.create_user(username="other-admin", role="admin")
        self.client.force_login(other)
        self.assertEqual(self.client.post(url, {"action": "apply", "preview_token": token}).status_code, 400)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.post(url, {"action": "apply", "preview_token": token}).status_code, 302)
        self.assertEqual(RundownImport.objects.count(), 1)
        viewer = User.objects.create_user(username="read-only", role="viewer")
        self.client.force_login(viewer)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertEqual(Client().get(url).status_code, 302)
        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.admin)
        self.assertEqual(csrf_client.post(url, {"action": "apply", "preview_token": token}).status_code, 403)

    def test_modal_preview_apply_and_errors_stay_in_dashboard(self):
        url = reverse("core:import_rundown")
        headers = {"HTTP_X_RUNDOWN_MODAL": "1"}
        response = self.client.get(url, **headers)
        self.assertContains(response, "data-rundown-import-content")
        self.assertNotContains(response, "<!DOCTYPE", html=False)
        self.assertNotContains(response, "Back to dashboard")
        response = self.client.post(url, {}, **headers)
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "Choose the exported", status_code=400)
        content = self.edit(lambda book: setattr(book["Piping - Fabrication"]["D10"], "value", 2))
        response = self.client.post(url, {"rundown_file": SimpleUploadedFile("updated.xlsx", content)}, **headers)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(RundownImport.objects.exists())
        response = self.client.post(url, {"action": "apply", "preview_token": response.context["preview_token"]}, **headers)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])
        payload = response.json()["payloads"]["fabrication:piping"]
        self.assertEqual(payload["kpis"]["actual_progress_pct"], 20)
        self.assertEqual(response["Cache-Control"], "no-store")
        self.assertEqual(RundownImport.objects.count(), 1)
        self.assertContains(self.client.get(url, **headers), "updated.xlsx")

    def test_zip_limits_report_size_and_entry_count_separately(self):
        from zipfile import ZipFile, ZIP_DEFLATED
        for extra_count, extra_size, message in [(201, 1, "internal files"), (1, 51 * 1024 * 1024, "limit: 50 MB")]:
            stream = BytesIO(self.content)
            with ZipFile(stream, "a", compression=ZIP_DEFLATED) as archive:
                for index in range(extra_count):
                    archive.writestr(f"extra{index}.txt", b" " * extra_size)
            with self.assertRaisesRegex(ValueError, message):
                wb.parse_workbook(stream.getvalue(), self.current)
        self.assertFalse(RundownImport.objects.exists())
