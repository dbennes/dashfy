from datetime import timedelta
from io import BytesIO
from uuid import uuid4

from django.test import TestCase, Client, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User, Client as Tenant
from apps.core.models import SectionNote, SectionNoteEvent


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class SectionNotesTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="ana", first_name="Ana", last_name="Silva", role="analyst")
        self.other = User.objects.create_user(username="bruno", role="viewer")
        self.client.force_login(self.user)
        self.url = reverse("core:section_notes")
        self.data = {"section": "s03", "body": "Material aguardando confirmação.",
            "date": timezone.localdate().isoformat(), "due_date": (timezone.localdate()+timedelta(days=7)).isoformat(),
            "information_only": False, "request_id": str(uuid4()), "author": "Forged identity"}

    def create(self, **changes):
        return self.client.post(self.url, dict(self.data, **changes), content_type="application/json")

    def status(self, note, value, version):
        return self.client.post(reverse("core:section_note_status", args=[note]),
                                {"status": value, "version": version}, content_type="application/json")

    def test_creation_records_server_identity_context_and_idempotency(self):
        response = self.create(context={"discipline": "piping", "secret": "must not persist"})
        self.assertEqual(response.status_code, 201)
        note = response.json()["note"]
        self.assertEqual(note["author"], "Ana Silva")
        self.assertEqual(note["status"], "pending")
        self.assertEqual(note["context"], {"discipline": "piping"})
        self.assertEqual(len(note["events"]), 1)
        self.assertEqual(note["events"][0]["actor"], "Ana Silva")
        self.assertEqual(self.create().status_code, 200)
        self.assertEqual(SectionNote.objects.count(), 1)
        self.assertEqual(SectionNoteEvent.objects.count(), 1)
        self.assertEqual(self.create(body="Different content").status_code, 409)

    def test_status_audit_cancel_restore_and_concurrent_update(self):
        note = self.create().json()["note"]
        self.client.force_login(self.other)
        resolved = self.status(note["id"], "resolved", 1)
        self.assertEqual(resolved.status_code, 200)
        self.assertEqual(resolved.json()["note"]["events"][-1]["actor"], "bruno")
        self.assertEqual(self.status(note["id"], "cancelled", 1).status_code, 409)
        cancelled = self.status(note["id"], "cancelled", 2)
        self.assertEqual(cancelled.json()["note"]["status"], "cancelled")
        restored = self.status(note["id"], "pending", 3)
        self.assertEqual(restored.json()["note"]["version"], 4)
        self.assertEqual(self.status(note["id"], "pending", 4).status_code, 200)
        self.assertEqual(SectionNoteEvent.objects.count(), 4)
        self.assertEqual(SectionNote.objects.get().body, self.data["body"])

    def test_information_only_dates_and_validation(self):
        self.assertEqual(self.create(due_date=None).status_code, 400)
        self.assertEqual(self.create(body=" ").status_code, 400)
        self.assertEqual(self.create(body="x"*4001).status_code, 400)
        self.assertEqual(self.create(section="unknown").status_code, 400)
        self.assertEqual(self.create(date=(timezone.localdate()+timedelta(days=1)).isoformat()).status_code, 400)
        self.assertEqual(self.create(due_date=(timezone.localdate()-timedelta(days=1)).isoformat()).status_code, 400)
        note = self.create(information_only=True, due_date=None).json()["note"]
        self.assertIsNone(note["due_date"])
        summary = self.client.get(self.url).json()["sections"]["s03"]
        self.assertEqual(summary["pending"], 0)
        self.assertEqual(summary["total"], 1)

    def test_tenant_isolation_applies_to_list_status_summary_and_export(self):
        note = self.create().json()["note"]
        tenant = Tenant.objects.create(name="Other tenant", slug="other")
        client_user = User.objects.create_user(username="client", role="client", client=tenant)
        self.client.force_login(client_user)
        self.assertEqual(self.client.get(self.url, {"section":"s03"}).json()["total"], 0)
        self.assertEqual(self.client.get(self.url).json()["sections"]["s03"]["total"], 0)
        self.assertEqual(self.status(note["id"], "resolved", 1).status_code, 404)
        exported = self.client.get(reverse("core:section_minutes"))
        book = load_workbook(BytesIO(exported.content))
        self.assertNotIn(self.data["body"], str(list(book.active.values)))
        book.close()

    def test_summary_deduplicates_people_and_history_filters_paginate(self):
        self.create()
        self.create(request_id=str(uuid4()), section="s01")
        for index in range(21):
            self.create(request_id=str(uuid4()), body=f"Record {index}")
        summary = self.client.get(self.url).json()["sections"]
        self.assertEqual(summary["s03"]["total"], 22)
        self.assertEqual(len(summary["s03"]["people"]), 1)
        first = self.client.get(self.url, {"section":"s03"}).json()
        second = self.client.get(self.url, {"section":"s03", "page":2}).json()
        self.assertEqual(len(first["notes"]), 20)
        self.assertTrue(first["has_next"])
        self.assertEqual(len(second["notes"]), 2)
        self.assertFalse(second["has_next"])
        self.assertEqual(self.client.get(self.url, {"section":"s03", "status":"resolved"}).json()["total"], 0)

    def test_single_sheet_minutes_include_events_and_never_execute_text(self):
        note = self.create(body='=HYPERLINK("https://example.invalid","note")').json()["note"]
        self.status(note["id"], "resolved", 1)
        self.status(note["id"], "cancelled", 2)
        response = self.client.get(reverse("core:section_minutes"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Cache-Control"], "no-store")
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Follow-up minutes"])
        sheet = book.active
        self.assertEqual(sheet.max_row, 8)
        self.assertEqual(sheet["F6"].data_type, "s")
        self.assertTrue(sheet["F6"].font.strike)
        self.assertEqual([sheet.cell(row,13).value for row in (6,7,8)], ["Pending", "Resolved", "Cancelled"])
        book.close()

    def test_authentication_and_csrf_are_required(self):
        self.assertEqual(Client().get(self.url).status_code, 302)
        self.assertEqual(Client().get(reverse("core:section_minutes")).status_code, 302)
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        self.assertEqual(client.post(self.url, self.data, content_type="application/json").status_code, 403)

    def test_panel_comments_and_authors_are_isolated_from_sibling_and_legacy(self):
        legacy = self.create().json()["note"]
        rundown = self.create(panel="piping-rundown", request_id=str(uuid4())).json()["note"]
        self.client.force_login(self.other)
        skyline = self.create(panel="wooden-box-skyline", request_id=str(uuid4())).json()["note"]
        self.assertEqual(rundown["panel"], "piping-rundown")
        self.assertEqual(skyline["panel"], "wooden-box-skyline")
        for key, expected in [("piping-rundown",rundown), ("wooden-box-skyline",skyline)]:
            records = self.client.get(self.url, {"panel":key}).json()["notes"]
            self.assertEqual([note["id"] for note in records], [expected["id"]])
        self.assertEqual([n["id"] for n in self.client.get(self.url, {"section":"s03"}).json()["notes"]], [legacy["id"]])
        self.status(rundown["id"], "resolved", 1)
        summary = self.client.get(self.url).json()["sections"]
        self.assertEqual(summary["piping-rundown"]["pending"], 0)
        self.assertEqual(summary["wooden-box-skyline"]["pending"], 1)
        self.assertEqual(summary["s03"]["total"], 1)
        self.assertEqual(summary["piping-rundown"]["people"][0]["name"], "Ana Silva")
        self.assertEqual(summary["wooden-box-skyline"]["people"][0]["name"], "bruno")
        self.assertEqual(self.create(panel="nonexistent", request_id=str(uuid4())).status_code, 400)
        self.assertEqual(self.create(panel="vessel-map", request_id=str(uuid4())).status_code, 400)
        vessel = self.create(panel="vessel-map", section="ais", request_id=str(uuid4())).json()["note"]
        self.assertEqual(vessel["panel"], "vessel-map")
        book = load_workbook(BytesIO(self.client.get(reverse("core:section_minutes")).content))
        labels = [book.active.cell(row,2).value for row in range(6,book.active.max_row+1)]
        self.assertIn("Fabrication / Piping ISO rundown", labels)
        self.assertIn("Fabrication / Wooden Box skyline", labels)
        book.close()

    def test_each_registered_panel_has_one_explicit_template_anchor(self):
        import re
        from pathlib import Path
        from django.conf import settings
        from apps.core.section_notes import PANELS
        html = "\n".join((Path(settings.BASE_DIR)/name).read_text(encoding="utf-8") for name in
                         ("templates/core/home.html", "templates/vessels/_dashboard.html"))
        keys = re.findall(r'data-note-panel="([^"]+)"', html)
        self.assertCountEqual(keys, PANELS)
        self.assertEqual(len(keys), len(set(keys)))
