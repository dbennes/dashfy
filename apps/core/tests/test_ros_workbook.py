from copy import deepcopy
from datetime import date
from hashlib import sha256
from io import BytesIO
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook

from apps.core import ros_workbook as ros
from apps.core.models import RosScheduleImport


class RosWorkbookTests(SimpleTestCase):
    def setUp(self):
        self.current = ros._default_schedule()
        self.content = ros.export_ros_workbook(self.current)

    def edited(self, edit):
        book = load_workbook(BytesIO(self.content))
        edit(book)
        output = BytesIO()
        book.save(output)
        book.close()
        return output.getvalue()

    def test_unchanged_export_preserves_exact_rows_curves_metadata_and_revision(self):
        parsed = ros.parse_ros_workbook(self.content, self.current)
        self.assertEqual(parsed["skyline"], self.current["skyline"])
        self.assertEqual(parsed["rundown"], self.current["rundown"])
        self.assertEqual(parsed["revision"], self.current["revision"])
        self.assertEqual(parsed["changes"], [])
        self.assertEqual(parsed["changed_rows"], 0)
        self.assertEqual(sum(row[3] for row in parsed["skyline"]["rows"]), 607)

    def test_reordering_rows_retains_stable_identity_without_data_changes(self):
        def edit(book):
            sheet = book[ros.SCHEDULE_SHEET]
            values = list(sheet.values)[1:]
            sheet.delete_rows(2, len(values))
            for row in reversed(values):
                sheet.append([row[0], row[1], row[2].date().isoformat(), row[3].date().isoformat(), row[4]])
        parsed = ros.parse_ros_workbook(self.edited(edit), self.current)
        self.assertEqual(parsed["changes"], [])
        self.assertEqual(parsed["skyline"], self.current["skyline"])

    def test_changed_date_and_spools_reconcile_both_series_and_extend_calendar(self):
        before = deepcopy(self.current)
        def edit(book):
            sheet = book[ros.SCHEDULE_SHEET]
            sheet["C2"] = "2027-01-03"
            sheet["D2"] = "2027-01-06"
            sheet["E2"] = 11
        parsed = ros.parse_ros_workbook(self.edited(edit), self.current)
        self.assertEqual(self.current, before)
        self.assertEqual(parsed["changed_rows"], 1)
        self.assertEqual(len(parsed["changes"]), 3)
        self.assertEqual(parsed["skyline"]["rows"][0][1:], ["2027-01-03", "2027-01-06", 11])
        chart = parsed["rundown"]
        self.assertTrue(set(before["rundown"]["dates"]).issubset(chart["dates"]))
        for value in ("2027-01-03", "2027-01-04", "2027-01-06", "2027-01-07"):
            self.assertIn(value, chart["dates"])
        self.assertEqual(chart["baseline_rundown"][0], 611)
        for scenario in ("baseline", "lookahead"):
            daily, remaining = chart[f"{scenario}_total"], chart[f"{scenario}_rundown"]
            self.assertEqual(sum(value or 0 for value in daily), 611)
            zero_index = remaining.index(0)
            for index in range(1, zero_index + 1):
                self.assertEqual(remaining[index], remaining[index - 1] - daily[index - 1])
            self.assertTrue(all(value is None for value in remaining[zero_index + 1:]))
            self.assertTrue(all(value is None for value in daily[zero_index + 1:]))

    def test_earlier_release_is_included_and_does_not_drop_original_grid(self):
        content = self.edited(lambda book: setattr(book[ros.SCHEDULE_SHEET]["D2"], "value", "2026-01-01"))
        parsed = ros.parse_ros_workbook(content, self.current)
        self.assertEqual(parsed["rundown"]["dates"][0], "2026-01-01")
        self.assertEqual(parsed["rundown"]["lookahead_total"][0], 7)
        self.assertEqual(parsed["rundown"]["lookahead_rundown"][0], 607)

    def test_only_report_date_changes_source_metadata_without_changing_chart(self):
        content = self.edited(lambda book: setattr(book["Settings"]["B5"], "value", date(2026, 9, 11)))
        parsed = ros.parse_ros_workbook(content, self.current)
        self.assertEqual(parsed["snapshot_date"], "2026-09-11")
        self.assertEqual(parsed["changed_rows"], 0)
        self.assertEqual(len(parsed["changes"]), 1)
        self.assertEqual(parsed["skyline"]["rows"], self.current["skyline"]["rows"])
        for key in ("dates", "baseline_total", "baseline_rundown", "lookahead_total", "lookahead_rundown"):
            self.assertEqual(parsed["rundown"][key], self.current["rundown"][key])

    def test_readonly_rundown_values_do_not_update_operational_schedule(self):
        def edit(book):
            book["Rundown"]["B2"] = 999999
            book["Rundown"]["C2"] = "=1+1"
        parsed = ros.parse_ros_workbook(self.edited(edit), self.current)
        self.assertEqual(parsed["changes"], [])
        self.assertEqual(parsed["rundown"], self.current["rundown"])

    def test_missing_duplicate_unknown_and_renamed_rows_are_rejected(self):
        edits = [
            lambda book: book[ros.SCHEDULE_SHEET].delete_rows(2),
            lambda book: setattr(book[ros.SCHEDULE_SHEET]["A2"], "value", "ROS-00002"),
            lambda book: setattr(book[ros.SCHEDULE_SHEET]["A2"], "value", "ROS-99999"),
            lambda book: setattr(book[ros.SCHEDULE_SHEET]["B2"], "value", "renamed line"),
        ]
        for edit in edits:
            with self.subTest(edit=edit), self.assertRaises(ros.RosWorkbookError):
                ros.parse_ros_workbook(self.edited(edit), self.current)

    def test_split_line_requires_one_consistent_baseline(self):
        rows = self.current["skyline"]["rows"]
        index = next(index for index, row in enumerate(rows) if sum(item[0] == row[0] for item in rows) > 1)
        content = self.edited(lambda book: setattr(book[ros.SCHEDULE_SHEET].cell(index + 2, 3), "value", "2028-01-01"))
        with self.assertRaisesRegex(ros.RosWorkbookError, "inconsistent baseline"):
            ros.parse_ros_workbook(content, self.current)

    def test_formulas_in_editable_dates_quantities_or_settings_are_rejected(self):
        for sheet, address in [(ros.SCHEDULE_SHEET, "C2"), (ros.SCHEDULE_SHEET, "E2"), ("Settings", "B5")]:
            with self.subTest(sheet=sheet, address=address):
                content = self.edited(lambda book: setattr(book[sheet][address], "value", "=1+1"))
                with self.assertRaisesRegex(ros.RosWorkbookError, "formulas"):
                    ros.parse_ros_workbook(content, self.current)

    def test_ambiguous_missing_and_invalid_dates_are_rejected(self):
        for value in ("09/11/2026", "2026-02-30", "20260911", "2026-09-11T12:00:00", None):
            with self.subTest(value=value):
                content = self.edited(lambda book: setattr(book[ros.SCHEDULE_SHEET]["C2"], "value", value))
                with self.assertRaisesRegex(ros.RosWorkbookError, "date"):
                    ros.parse_ros_workbook(content, self.current)

    def test_fractional_zero_negative_text_and_boolean_quantities_are_rejected(self):
        for value in (0, -1, 1.5, "7", True, None):
            with self.subTest(value=value):
                content = self.edited(lambda book: setattr(book[ros.SCHEDULE_SHEET]["E2"], "value", value))
                with self.assertRaisesRegex(ros.RosWorkbookError, "spools"):
                    ros.parse_ros_workbook(content, self.current)

    def test_changed_headers_format_and_stale_revision_are_rejected(self):
        for sheet, address, value in [(ros.SCHEDULE_SHEET, "C1", "New date"), ("Settings", "B2", "Another workbook"), ("Settings", "B3", 2), ("Settings", "B4", "old-revision")]:
            with self.subTest(address=address):
                content = self.edited(lambda book: setattr(book[sheet][address], "value", value))
                with self.assertRaises(ros.RosWorkbookError):
                    ros.parse_ros_workbook(content, self.current)

    def test_bad_zip_file_size_and_expanded_size_are_rejected(self):
        for content in (b"not an Excel workbook", b""):
            with self.assertRaises(ros.RosWorkbookError):
                ros.parse_ros_workbook(content, self.current)
        with patch.object(ros, "MAX_FILE_BYTES", 1), self.assertRaisesRegex(ros.RosWorkbookError, "10 MB"):
            ros.parse_ros_workbook(self.content, self.current)
        with patch.object(ros, "MAX_EXPANDED_BYTES", 1), self.assertRaisesRegex(ros.RosWorkbookError, "expanded"):
            ros.parse_ros_workbook(self.content, self.current)

    def test_dimension_limits_and_undeclared_extra_columns_are_rejected(self):
        content = self.edited(lambda book: setattr(book[ros.SCHEDULE_SHEET]["F2"], "value", "extra"))
        with self.assertRaisesRegex(ros.RosWorkbookError, "extra columns"):
            ros.parse_ros_workbook(content, self.current)
        with patch.object(ros, "MAX_ROWS", 175), self.assertRaisesRegex(ros.RosWorkbookError, "size limits"):
            ros.parse_ros_workbook(self.content, self.current)

    def test_understated_dimensions_cannot_hide_an_added_row(self):
        content = self.edited(lambda book: book[ros.SCHEDULE_SHEET].append(["ROS-99999", "untracked", "2026-09-10", "2026-09-10", 2]))
        rewritten = BytesIO()
        with ZipFile(BytesIO(content)) as source, ZipFile(rewritten, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet2.xml":
                    payload = payload.replace(b'<dimension ref="A1:E177"', b'<dimension ref="A1:E176"')
                target.writestr(item.filename, payload)
        with self.assertRaisesRegex(ros.RosWorkbookError, "unknown Row ID"):
            ros.parse_ros_workbook(rewritten.getvalue(), self.current)

    def test_truncated_worksheet_xml_has_a_friendly_import_error(self):
        rewritten = BytesIO()
        with ZipFile(BytesIO(self.content)) as source, ZipFile(rewritten, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet2.xml":
                    payload = payload[:-50]
                target.writestr(item.filename, payload)
        with self.assertRaisesRegex(ros.RosWorkbookError, "could not be read"):
            ros.parse_ros_workbook(rewritten.getvalue(), self.current)


class RosSchedulePersistenceTests(TestCase):
    def setUp(self):
        self.current = ros.load_current_schedule()
        self.user = get_user_model().objects.create_user(username="ros-editor")
        book = load_workbook(BytesIO(ros.export_ros_workbook(self.current)))
        book[ros.SCHEDULE_SHEET]["E2"] = 11
        output = BytesIO()
        book.save(output)
        book.close()
        self.content = output.getvalue()
        self.parsed = ros.parse_ros_workbook(self.content, self.current)

    def apply(self, parsed=None):
        return ros.apply_ros_import(parsed or self.parsed, "ros-update.xlsx", sha256(self.content).hexdigest(), len(self.content), self.user)

    def test_noop_does_not_create_an_import_or_change_source(self):
        parsed = ros.parse_ros_workbook(ros.export_ros_workbook(self.current), self.current)
        self.assertEqual(self.apply(parsed), (None, False))
        self.assertFalse(RosScheduleImport.objects.exists())
        self.assertEqual(ros.load_current_schedule(), self.current)

    def test_apply_persists_shared_snapshot_and_user_file_provenance(self):
        batch, created = self.apply()
        self.assertTrue(created)
        self.assertEqual(batch.imported_by, self.user)
        self.assertEqual(batch.original_filename, "ros-update.xlsx")
        self.assertEqual(batch.file_hash, sha256(self.content).hexdigest())
        self.assertEqual(batch.file_size, len(self.content))
        self.assertEqual(batch.base_revision, self.current["revision"])
        current = ros.load_current_schedule()
        self.assertEqual(current["revision"], str(batch.revision))
        self.assertEqual(current["skyline"]["rows"][0][3], 11)
        self.assertEqual(current["rundown"]["baseline_rundown"][0], 611)
        self.assertEqual(current["skyline"]["source"]["workbook"], "ros-update.xlsx")
        self.assertEqual(current["rundown"]["source"]["managed_by"], "ros_workbook")
        second = ros.parse_ros_workbook(ros.export_ros_workbook(current), current)
        self.assertEqual(second["changes"], [])
        self.assertEqual(second["skyline"], current["skyline"])
        same, created_again = self.apply(second)
        self.assertFalse(created_again)
        self.assertEqual(same.pk, batch.pk)
        self.assertEqual(RosScheduleImport.objects.count(), 1)

    def test_stale_preview_cannot_overwrite_newer_import(self):
        self.apply()
        with self.assertRaisesRegex(ros.RosWorkbookError, "changed after"):
            self.apply()
        self.assertEqual(RosScheduleImport.objects.count(), 1)
        self.assertEqual(ros.load_current_schedule()["skyline"]["rows"][0][3], 11)

    def test_only_one_import_can_be_based_on_each_revision(self):
        batch, _ = self.apply()
        with self.assertRaises(IntegrityError), transaction.atomic():
            RosScheduleImport.objects.create(
                base_revision=batch.base_revision, original_filename="competing.xlsx",
                snapshot_date=date(2026, 9, 2), payload=batch.payload,
            )
        self.assertEqual(RosScheduleImport.objects.count(), 1)

    def test_database_error_does_not_silently_restore_seed(self):
        from django.db import DatabaseError
        with patch.object(RosScheduleImport.objects, "order_by", side_effect=DatabaseError("unavailable")):
            with self.assertRaises(DatabaseError):
                ros.load_current_schedule()
