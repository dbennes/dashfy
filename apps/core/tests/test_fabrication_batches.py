from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import Workbook

from apps.core.fabrication_batches import parse_workbook, import_workbook, current_batches
from apps.core.models import FabricationBatchImport
from apps.accounts.models import User
from apps.core.release_scope import release_scope_context


def workbook_bytes():
    book = Workbook()
    sps = book.active
    sps.title = 'SPS BATCH 1-6'
    sps.append(['S/N', 'DRAWING', 'Line No.', 'Rev.', 'Page', 'CPMTO New', 'TRAVEL PACK', 'SPS Drawing', 'BATCH NO.', 'PRIORITY', 'T.P Status', 'Fabrication (Cutting/Fit-up/Welding)'])
    sps.append([1, 'DWG1', '14"-PM-123-CC', '', '', 'SPS-1', '', '', 'Batch 1 - B', '', '', 'OK'])
    sps.append([None, None, None, '', '', 'SPS-2', '', '', 'Batch 1 - B', '', '', 'ONGOING'])
    sps.merge_cells('B2:B3'); sps.merge_cells('C2:C3')
    sps.append([2, 'DWG2', 'LINE2', '', '', 'NA', '', '', 'Batch 2A', '', '', None])
    dummy = book.create_sheet('DUMMY SPOOLS')
    dummy.append(['SN', 'Valve Description', 'Line Number', 'Reference Drawing', 'TRAVEL PACK STATUS', 'FABRICATION', 'RFLO/LOADOUT SATUS'])
    dummy.append([1, 'VALVE1', '14-PM-123-CC', 'DWG1', '', 'OK'])
    dummy.append([2, 'VALVE2', '14-PM-123-CC', 'DWG1', '', None])
    dummy.append(['ANGULAR DUMMY SPOOLS'])
    dummy.append(['SN', 'Valve Description', 'Line Number', 'Reference Drawing'])
    dummy.append([1, 'VALVE3', 'LINE 2', 'DWG2', '', None])
    dummy.append([2, 'VALVE4', 'UNKNOWN', 'OTHER', '', 'OK'])
    stream = BytesIO(); book.save(stream); book.close()
    return stream.getvalue()


class BatchParsingTests(SimpleTestCase):
    def test_merges_exclusions_subbatches_and_dummy_mapping(self):
        result = parse_workbook(workbook_bytes())
        self.assertEqual(result['labels'], ['Batch 1B', 'Batch 2A', 'Unassigned'])
        self.assertEqual(result['scopes']['structural']['total'], [2, 0, 0])
        self.assertEqual(result['scopes']['structural']['done'], [1, 0, 0])
        self.assertEqual(result['scopes']['piping']['total'], [2, 1, 1])
        self.assertEqual(result['scopes']['piping']['done'], [1, 0, 1])
        self.assertEqual(result['excluded_sps_rows'], 1)
        self.assertEqual(result['unassigned_items'], 1)

    def test_invalid_workbook_is_rejected(self):
        with self.assertRaisesMessage(ValueError, 'valid .xlsx'):
            parse_workbook(b'invalid')


class BatchImportTests(TestCase):
    def test_latest_import_replaces_summary_without_accumulating(self):
        for _ in range(2):
            import_workbook(SimpleUploadedFile('batches.xlsx', workbook_bytes()))
        self.assertEqual(FabricationBatchImport.objects.count(), 2)
        self.assertEqual(current_batches()['payload']['scopes']['all']['item_count'], 6)
        with self.assertRaises(ValueError):
            import_workbook(SimpleUploadedFile('bad.xlsx', b'bad'))
        self.assertEqual(FabricationBatchImport.objects.count(), 2)

    def test_import_requires_admin_and_accepts_workbook(self):
        user = User.objects.create_user(username='batch-viewer', password='test')
        self.client.force_login(user)
        url = reverse('core:import_fabrication_batches')
        self.assertEqual(self.client.post(url).status_code, 403)
        user.role = 'admin'
        user.save()
        response = self.client.post(url, {'batch_file': SimpleUploadedFile('batch.xlsx', workbook_bytes())})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(FabricationBatchImport.objects.count(), 1)

    def test_preparing_import_is_separate_and_both_uses_only_latest_per_scope(self):
        for _ in range(2):
            official = import_workbook(SimpleUploadedFile('official.xlsx', workbook_bytes()))
        self.assertEqual(official.payload['release_scope'], 'official')
        with release_scope_context('preparing'):
            self.assertEqual(current_batches(), {})
            for _ in range(2):
                preparing = import_workbook(SimpleUploadedFile('preparing.xlsx', workbook_bytes()))
            preparing.payload['labels'][0] = 'Batch 3'
            preparing.save(update_fields=['payload'])
            self.assertEqual(preparing.payload['release_scope'], 'preparing')
            self.assertEqual(current_batches()['filename'], 'preparing.xlsx')
            self.assertEqual(current_batches()['payload']['scopes']['all']['item_count'], 6)
        self.assertEqual(current_batches()['filename'], 'official.xlsx')
        with release_scope_context('all'):
            result = current_batches()
            self.assertEqual(result['filename'], 'official.xlsx + preparing.xlsx')
            self.assertEqual(result['payload']['scopes']['all']['item_count'], 12)
            self.assertEqual(result['payload']['labels'], ['Batch 1B', 'Batch 2A', 'Batch 3', 'Unassigned'])
            self.assertEqual(result['payload']['scopes']['all']['total'], [4, 2, 4, 2])
            self.assertEqual(result['payload']['scopes']['all']['done'], [2, 0, 2, 2])
            self.assertEqual(result['payload']['unassigned_items'], 2)

    def test_legacy_batches_are_official_and_both_import_is_rejected(self):
        FabricationBatchImport.objects.create(original_filename='legacy.xlsx', file_hash='legacy',
                                               payload=parse_workbook(workbook_bytes()))
        self.assertEqual(current_batches()['filename'], 'legacy.xlsx')
        with release_scope_context('preparing'):
            self.assertEqual(current_batches(), {})
        with release_scope_context('all'):
            self.assertEqual(current_batches()['payload']['scopes']['all']['item_count'], 6)
            with self.assertRaisesMessage(ValueError, 'Choose Official or Preparing'):
                import_workbook(SimpleUploadedFile('both.xlsx', workbook_bytes()))
        self.assertEqual(FabricationBatchImport.objects.count(), 1)

    def test_batch_upload_posts_explicit_scope_and_redirect_keeps_selection(self):
        user = User.objects.create_user(username='batch-admin-scope', role='admin')
        self.client.force_login(user)
        url = reverse('core:import_fabrication_batches')
        response = self.client.post(url, {'release_scope': 'preparing',
                                         'batch_file': SimpleUploadedFile('preparing.xlsx', workbook_bytes())})
        self.assertEqual(response.status_code, 302)
        self.assertIn('?release_scope=preparing#s03', response.url)
        self.assertEqual(FabricationBatchImport.objects.get().payload['release_scope'], 'preparing')
        response = self.client.post(url, {'release_scope': 'all',
                                         'batch_file': SimpleUploadedFile('both.xlsx', workbook_bytes())})
        self.assertEqual(response.status_code, 302)
        self.assertIn('?release_scope=all#s03', response.url)
        self.assertEqual(FabricationBatchImport.objects.count(), 1)
        self.assertTrue(any('Choose Official or Preparing' in str(message)
                            for message in response.wsgi_request._messages))
