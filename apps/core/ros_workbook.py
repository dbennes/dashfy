"""Editable ROS workbooks; live DATAFY evidence is never written by this module."""
from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
from datetime import date, datetime, time, timedelta
from hashlib import sha256
from io import BytesIO
import json
import math
from pathlib import Path
from typing import Any
from uuid import uuid4
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile
from zlib import error as ZipCompressionError

from django.db import IntegrityError, transaction
from django.utils import timezone

from .models import RosScheduleImport


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 1000
MAX_ROWS = 10000
MAX_COLUMNS = 128
MAX_CELLS = 500000
WORKBOOK_FORMAT = "DASHFY ROS schedule"
WORKBOOK_VERSION = 1
SCHEDULE_SHEET = "ROS Schedule"
HEADERS = ("Row ID", "Line", "Baseline Date", "Lookahead Date", "Spools")
_SETTINGS_KEYS = ("Format", "Version", "Revision", "Data date")
_FIELDS = ("Baseline Date", "Lookahead Date", "Spools")


class RosWorkbookError(ValueError):
    """A rejected workbook or a preview that is no longer current."""


def _row_id(index: int) -> str:
    return f"ROS-{index + 1:05d}"


def _validate_payloads(skyline: dict, rundown: dict) -> None:
    # Function imports prevent a cycle when the display sources call this module.
    from .rundown_source import _validated_snapshot as validate_rundown
    from .skyline_source import _validated_snapshot as validate_skyline

    validate_skyline(skyline)
    validate_rundown(rundown)
    expected = sum(row[3] for row in skyline["rows"])
    if rundown["baseline_rundown"][0] != expected:
        raise RosWorkbookError("The ROS schedule and Piping rundown have different spool totals.")


def _default_schedule() -> dict[str, Any]:
    from .rundown_source import RUNDOWN_DATA_PATH
    from .skyline_source import SKYLINE_DATA_PATH

    skyline = json.loads(SKYLINE_DATA_PATH.read_text(encoding="utf-8"))
    rundown = json.loads(RUNDOWN_DATA_PATH.read_text(encoding="utf-8"))
    _validate_payloads(skyline, rundown)
    canonical = json.dumps([skyline, rundown], sort_keys=True, separators=(",", ":"))
    return {
        "revision": sha256(canonical.encode("utf-8")).hexdigest(),
        "skyline": skyline,
        "rundown": rundown,
        "batch": None,
    }


def _batch_schedule(batch: RosScheduleImport) -> dict[str, Any]:
    skyline = deepcopy(batch.payload["skyline"])
    rundown = deepcopy(batch.payload["rundown"])
    _validate_payloads(skyline, rundown)
    return {"revision": str(batch.revision), "skyline": skyline, "rundown": rundown, "batch": batch}


def load_current_schedule() -> dict[str, Any]:
    """Load the last accepted revision, or repository seed if none was imported.

    Database errors deliberately propagate; showing an old seed after an outage
    would silently replace the user's accepted operational schedule.
    """
    batch = RosScheduleImport.objects.order_by("-pk").first()
    return _batch_schedule(batch) if batch else _default_schedule()


def export_ros_workbook(current: dict[str, Any]) -> bytes:
    import xlsxwriter

    _validate_payloads(current["skyline"], current["rundown"])
    output = BytesIO()
    book = xlsxwriter.Workbook(output, {
        "in_memory": True, "strings_to_urls": False, "strings_to_formulas": False,
    })
    header = book.add_format({"bold": True, "bg_color": "#183D5D", "font_color": "#FFFFFF", "border": 1})
    locked = book.add_format({"locked": True, "bg_color": "#EDF1F5"})
    editable_date = book.add_format({"locked": False, "num_format": "yyyy-mm-dd", "bg_color": "#FFF3CC"})
    editable_number = book.add_format({"locked": False, "num_format": "0", "bg_color": "#FFF3CC"})
    note = book.add_format({"text_wrap": True, "valign": "top"})

    instructions = book.add_worksheet("Instructions")
    instructions.set_column("A:A", 110)
    instructions.write(0, 0, "ROS schedule: export, edit and import", header)
    for index, text in enumerate([
        "Edit yellow cells in ROS Schedule: Baseline Date, Lookahead Date and Spools. Dates use YYYY-MM-DD. Spools must be positive whole numbers.",
        "Update Data date in Settings to identify the source report date. It does not change the dashboard's current-date cutoff.",
        "Keep every Row ID and Line unchanged. Do not add or delete rows. Reordering rows is allowed. Split rows for the same line must share one baseline date.",
        "Upload this workbook using Import ROS. Review the changed dates, quantities and source date, then apply the preview.",
        "Rundown is calculated from ROS Schedule and shown for reference. Changes made to Rundown are not imported.",
        "DATAFY material lights, AVEON fabrication dates and Structural rundown update from their source systems. They are not editable in this workbook.",
        "Only the latest exported revision can be imported. If another person updates ROS, export again and copy your edits into the new workbook.",
        "The sheet protection prevents accidental edits; it is not a security control. The importer validates row identities, dates and quantities again.",
    ], 2):
        instructions.write(index, 0, text, note)
        instructions.set_row(index, 40)
    instructions.protect()

    schedule = book.add_worksheet(SCHEDULE_SHEET)
    schedule.freeze_panes(1, 2)
    schedule.write_row(0, 0, HEADERS, header)
    schedule.set_column("A:A", 15)
    schedule.set_column("B:B", 32)
    schedule.set_column("C:D", 20)
    schedule.set_column("E:E", 12)
    for index, row in enumerate(current["skyline"]["rows"], 1):
        schedule.write_string(index, 0, _row_id(index - 1), locked)
        schedule.write_string(index, 1, row[0], locked)
        schedule.write_datetime(index, 2, datetime.combine(date.fromisoformat(row[1]), time()), editable_date)
        schedule.write_datetime(index, 3, datetime.combine(date.fromisoformat(row[2]), time()), editable_date)
        schedule.write_number(index, 4, row[3], editable_number)
    last = len(current["skyline"]["rows"])
    schedule.autofilter(0, 0, last, 4)
    schedule.data_validation(1, 4, last, 4, {
        "validate": "integer", "criteria": ">", "value": 0,
        "error_title": "Invalid spool quantity", "error_message": "Enter a positive whole number.",
    })
    schedule.protect(options={"autofilter": True, "select_locked_cells": True, "select_unlocked_cells": True})

    settings = book.add_worksheet("Settings")
    settings.set_column("A:A", 22)
    settings.set_column("B:B", 72)
    settings.write_row(0, 0, ["Setting", "Value"], header)
    for index, (key, value) in enumerate(zip(_SETTINGS_KEYS[:3], [WORKBOOK_FORMAT, WORKBOOK_VERSION, current["revision"]]), 1):
        settings.write_string(index, 0, key, locked)
        settings.write(index, 1, value, locked)
    settings.write_string(4, 0, "Data date", locked)
    settings.write_datetime(4, 1, datetime.combine(date.fromisoformat(current["skyline"]["source"]["snapshot_date"]), time()), editable_date)
    settings.protect()

    rundown = book.add_worksheet("Rundown")
    rundown.freeze_panes(1, 1)
    rundown.set_column("A:A", 16)
    rundown.set_column("B:E", 22)
    rundown.write_row(0, 0, ["Date", "Baseline releases", "Baseline remaining", "Lookahead releases", "Lookahead remaining"], header)
    raw = current["rundown"]
    keys = ("baseline_total", "baseline_rundown", "lookahead_total", "lookahead_rundown")
    for index, value in enumerate(raw["dates"], 1):
        rundown.write_string(index, 0, value, locked)
        for column, key in enumerate(keys, 1):
            rundown.write(index, column, raw[key][index - 1], locked)
    rundown.protect()
    book.close()
    return output.getvalue()


def _workbook_date(value: Any, label: str) -> str:
    if isinstance(value, datetime):
        if value.time() != time():
            raise RosWorkbookError(f"{label}: use a date without a time.")
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        candidate = value.strip()
        try:
            parsed = date.fromisoformat(candidate)
        except ValueError:
            pass
        else:
            if candidate == parsed.isoformat():
                return candidate
    raise RosWorkbookError(f"{label}: enter a valid date as YYYY-MM-DD or an Excel date.")


def _spools(value: Any, label: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise RosWorkbookError(f"{label}: spools must be a positive whole number.")
    if not math.isfinite(value) or value <= 0 or not float(value).is_integer() or value > 2147483647:
        raise RosWorkbookError(f"{label}: spools must be a positive whole number below 2,147,483,648.")
    return int(value)


def _guard_archive(content: bytes) -> None:
    if not content or len(content) > MAX_FILE_BYTES:
        raise RosWorkbookError("Upload an Excel workbook no larger than 10 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                raise RosWorkbookError("The workbook exceeds the expanded-file size limit.")
            if any(item.flag_bits & 1 for item in entries):
                raise RosWorkbookError("Password-encrypted workbooks are not supported.")
            names = {item.filename for item in entries}
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise RosWorkbookError("Upload a valid .xlsx ROS workbook exported by DASHFY.")
    except BadZipFile as exc:
        raise RosWorkbookError("Upload a valid .xlsx ROS workbook exported by DASHFY.") from exc


def _read_sheet(sheet) -> list[list[Any]]:
    # Check declared dimensions before iteration, then reset them so forged or
    # stale Excel dimensions cannot hide populated cells from the validation.
    declared_rows, declared_columns = sheet.max_row or 0, sheet.max_column or 0
    if declared_rows > MAX_ROWS or declared_columns > MAX_COLUMNS or declared_rows * declared_columns > MAX_CELLS:
        raise RosWorkbookError(f'Worksheet "{sheet.title}" exceeds the workbook size limits.')
    sheet.reset_dimensions()
    result = []
    cells_seen = 0
    for row_number, row in enumerate(sheet.iter_rows(), 1):
        cells_seen += len(row)
        if row_number > MAX_ROWS or len(row) > MAX_COLUMNS or cells_seen > MAX_CELLS:
            raise RosWorkbookError(f'Worksheet "{sheet.title}" exceeds the workbook size limits.')
        if any(cell.data_type == "f" for cell in row):
            raise RosWorkbookError(f'{sheet.title} row {row_number}: formulas are not accepted; enter the value directly.')
        result.append([cell.value for cell in row])
    return result


def _derived_rundown(skyline: dict, current: dict) -> dict:
    raw = deepcopy(current)
    dates = {date.fromisoformat(value) for value in current["dates"]}
    releases = {"baseline": defaultdict(int), "lookahead": defaultdict(int)}
    for row in skyline["rows"]:
        for scenario, position in (("baseline", 1), ("lookahead", 2)):
            point = date.fromisoformat(row[position])
            dates.add(point)
            releases[scenario][point] += row[3]
    try:
        for daily in releases.values():
            dates.add(max(daily) + timedelta(days=1))
    except OverflowError as exc:
        raise RosWorkbookError("The last release date must allow a following day for the rundown balance.") from exc
    points = sorted(dates)
    scope = sum(row[3] for row in skyline["rows"])
    raw["dates"] = [point.isoformat() for point in points]
    for scenario, daily in releases.items():
        remaining = scope
        finished = False
        totals, balances = [], []
        for point in points:
            totals.append(None if finished else daily.get(point, 0))
            balances.append(None if finished else remaining)
            if remaining == 0:
                finished = True
            remaining -= daily.get(point, 0)
        raw[f"{scenario}_total"] = totals
        raw[f"{scenario}_rundown"] = balances
    return raw


def parse_ros_workbook(content: bytes, current: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import load_workbook

    _guard_archive(content)
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise RosWorkbookError("The Excel workbook could not be read. Export a new ROS workbook and try again.") from exc
    try:
        if "Settings" not in book.sheetnames or SCHEDULE_SHEET not in book.sheetnames:
            raise RosWorkbookError("This is not a ROS workbook. Export ROS from DASHFY before editing it.")
        setting_rows = _read_sheet(book["Settings"])
        if not setting_rows or setting_rows[0] != ["Setting", "Value"]:
            raise RosWorkbookError("Settings headers were changed. Export a new ROS workbook.")
        settings = {}
        for row in setting_rows[1:]:
            if not any(value is not None for value in row):
                continue
            if len(row) != 2 or row[0] not in _SETTINGS_KEYS or row[0] in settings:
                raise RosWorkbookError("Settings were changed or duplicated. Keep the exported settings and edit only Data date.")
            settings[row[0]] = row[1]
        if set(settings) != set(_SETTINGS_KEYS) or settings["Format"] != WORKBOOK_FORMAT or settings["Version"] != WORKBOOK_VERSION:
            raise RosWorkbookError("Unsupported ROS workbook format. Export a new ROS workbook.")
        if settings["Revision"] != current["revision"]:
            raise RosWorkbookError("This ROS workbook is out of date. Export the latest revision and copy your edits into it.")
        snapshot_date = _workbook_date(settings["Data date"], "Settings / Data date")
        schedule_rows = _read_sheet(book[SCHEDULE_SHEET])
        if not schedule_rows or tuple(schedule_rows[0]) != HEADERS:
            raise RosWorkbookError("ROS Schedule headers were changed. Keep the five exported columns in their original order.")
        expected = {_row_id(index): row for index, row in enumerate(current["skyline"]["rows"])}
        edited = {}
        for row_number, values in enumerate(schedule_rows[1:], 2):
            if not any(value is not None for value in values):
                continue
            values = values + [None] * (5 - len(values))
            if len(values) != 5 or values[0] not in expected:
                raise RosWorkbookError(f"ROS Schedule row {row_number}: unknown Row ID or extra columns. Do not add rows or columns.")
            row_id, line, baseline, lookahead, quantity = values
            if row_id in edited:
                raise RosWorkbookError(f"ROS Schedule row {row_number}: duplicate Row ID {row_id}.")
            if line != expected[row_id][0]:
                raise RosWorkbookError(f"ROS Schedule row {row_number}: Line does not match {row_id}. Keep Row ID and Line unchanged.")
            edited[row_id] = [
                line,
                _workbook_date(baseline, f"ROS Schedule row {row_number} / Baseline Date"),
                _workbook_date(lookahead, f"ROS Schedule row {row_number} / Lookahead Date"),
                _spools(quantity, f"ROS Schedule row {row_number}"),
            ]
        missing = set(expected) - set(edited)
        if missing:
            raise RosWorkbookError(f"The workbook is missing {len(missing)} ROS row(s). Keep every exported row; rows cannot be deleted.")
    except RosWorkbookError:
        raise
    except (ParseError, BadZipFile, ZipCompressionError, EOFError, ValueError) as exc:
        raise RosWorkbookError("The Excel worksheet could not be read. Export a new ROS workbook and try again.") from exc
    finally:
        book.close()

    skyline = deepcopy(current["skyline"])
    skyline["rows"] = [edited[row_id] for row_id in expected]
    changes = []
    changed_ids = set()
    for row_id, before in expected.items():
        after = edited[row_id]
        for position, field in enumerate(_FIELDS, 1):
            if before[position] != after[position]:
                changed_ids.add(row_id)
                changes.append({"row_id": row_id, "line": before[0], "field": field, "before": before[position], "after": after[position]})
    if snapshot_date != current["skyline"]["source"]["snapshot_date"]:
        changes.append({"row_id": "settings", "line": "Source report", "field": "Data date", "before": current["skyline"]["source"]["snapshot_date"], "after": snapshot_date})
    try:
        # Validate line-level baseline consistency before deriving either view.
        from .skyline_source import _validated_snapshot
        _validated_snapshot(skyline)
        rundown = _derived_rundown(skyline, current["rundown"]) if changed_ids else deepcopy(current["rundown"])
        if changes:
            skyline["source"]["snapshot_date"] = snapshot_date
            rundown["source"]["snapshot_date"] = snapshot_date
        _validate_payloads(skyline, rundown)
    except ValueError as exc:
        raise RosWorkbookError(str(exc)) from exc
    return {
        "revision": current["revision"], "skyline": skyline, "rundown": rundown,
        "changes": changes, "changed_rows": len(changed_ids), "snapshot_date": snapshot_date,
    }


def apply_ros_import(parsed: dict[str, Any], filename: str, file_hash: str, file_size: int, user) -> tuple[RosScheduleImport | None, bool]:
    """Apply one reviewed revision atomically; stale previews never overwrite it."""
    try:
        with transaction.atomic():
            previous = RosScheduleImport.objects.select_for_update().order_by("-pk").first()
            current = _batch_schedule(previous) if previous else _default_schedule()
            if parsed["revision"] != current["revision"]:
                raise RosWorkbookError("ROS changed after this workbook was exported. Export the latest revision and review your edits again.")
            if not parsed["changes"]:
                return previous, False
            skyline, rundown = deepcopy(parsed["skyline"]), deepcopy(parsed["rundown"])
            _validate_payloads(skyline, rundown)
            revision = uuid4()
            filename = Path(filename.replace("\\", "/")).name[:255]
            timestamp = timezone.now().isoformat()
            for raw in (skyline, rundown):
                raw["source"].update({
                    "workbook": filename, "workbook_sha256": file_hash,
                    "revision_timestamp": timestamp, "snapshot_date": parsed["snapshot_date"],
                    "import_revision": str(revision), "managed_by": "ros_workbook",
                })
            skyline["source"].update({
                "worksheet": SCHEDULE_SHEET, "range": f"A1:E{len(skyline['rows']) + 1}",
                "forecast_scope": "ROS Schedule / Baseline Date", "lookahead_scope": "ROS Schedule / Lookahead Date",
                "normalization": "Manual ROS line dates and spool quantities from the accepted DASHFY workbook. Row identity and split-line baseline dates are validated on import.",
            })
            rundown["source"].update({
                "worksheet": "Rundown", "range": f"A1:E{len(rundown['dates']) + 1}",
                "reconciled_from": "ROS Schedule / dates and spools; remaining is the start-of-day balance",
            })
            batch = RosScheduleImport.objects.create(
                revision=revision, base_revision=current["revision"], original_filename=filename,
                file_hash=file_hash, file_size=file_size,
                imported_by=user if user is not None and user.is_authenticated else None,
                snapshot_date=date.fromisoformat(parsed["snapshot_date"]),
                payload={"skyline": skyline, "rundown": rundown},
                metadata={"changes": deepcopy(parsed["changes"]), "changed_rows": parsed["changed_rows"], "row_count": len(skyline["rows"]), "scope_spools": sum(row[3] for row in skyline["rows"])},
            )
            return batch, True
    except IntegrityError as exc:
        if RosScheduleImport.objects.filter(base_revision=parsed["revision"]).exists():
            raise RosWorkbookError("ROS was updated by another import. Export the latest revision and review your edits again.") from exc
        raise
